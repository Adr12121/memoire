# -*- coding: utf-8 -*-
"""
Générateur maître de la soutenance PFE — Adrien TRAVAILLÉ
Diplôme d'Ingénieur Topographe INSA Strasbourg
Cabinet GEO-SIAPP (Aubenas, Ardèche)

Structure (35 diapositives) :
  Slide  1 : Page de titre (fond bleu nuit)
  Slide  2 : Sommaire visuel (8 tuiles)
  Slide  3 : Problématique
  Slide  4 : Transition Chapitre 1
  Slides  5- 6 : Chapitre 1 — Introduction
  Slide  7 : Transition Chapitre 2
  Slides  8-10 : Chapitre 2 — Analyse Métier
  Slide 11 : Transition Chapitre 3
  Slides 12-14 : Chapitre 3 — État de l'Art
  Slide 15 : Transition Chapitre 4
  Slides 16-19 : Chapitre 4 — Architecture & Pipeline
  Slide 20 : Transition Chapitre 5
  Slides 21-23 : Chapitre 5 — Validation (dont vidéo)
  Slide 24 : Transition Chapitre 6
  Slides 25-26 : Chapitre 6 — Intégration Géofoncier
  Slide 27 : Transition Chapitre 7
  Slide 28 : Chapitre 7 — Limites & Perspectives
  Slide 29 : Transition Chapitre 8
  Slides 30-31 : Chapitre 8 — Conclusion + Remerciements
  Slides 32-35 : Annexes jury (GLiNER, YOLOv8, OCR, Références)
"""

import os
from pptx import Presentation

from generate_deck_part1 import build_part1_presentation
from generate_deck_part2 import add_slides_part2
from generate_deck_part3 import add_slides_part3
from generate_deck_part4 import add_slides_part4


def generate_complete_defense_deck():
    print("=" * 70)
    print("DEMARRAGE DU GENERATEUR MAITRE — SOUTENANCE PFE (35 DIAPOSITIVES)")
    print("=" * 70)

    # Phase 1 : Slides 1 à 11 (titre, sommaire, problématique, chap. 1 & 2)
    print("-> Phase 1/4 : Slides 1 a 11 (Titre, Sommaire, Problematique, Chap. 1 & 2)...")
    prs = build_part1_presentation()
    print(f"   Slides generees : {len(prs.slides)}")

    # Phase 2 : Slides 12 à 15 (chap. 3 état de l'art + transition chap. 4)
    print("-> Phase 2/4 : Slides 12 a 15 (Chapitre 3 — Etat de l'art)...")
    add_slides_part2(prs)
    print(f"   Slides generees : {len(prs.slides)}")

    # Phase 3 : Slides 16 à 26 (chap. 4, 5 avec vidéo, 6)
    print("-> Phase 3/4 : Slides 16 a 26 (Chapitres 4, 5, 6)...")
    add_slides_part3(prs)
    print(f"   Slides generees : {len(prs.slides)}")

    # Phase 4 : Slides 27 à 35 (chap. 7, 8, annexes jury, références)
    print("-> Phase 4/4 : Slides 27 a 35 (Chapitres 7, 8, Annexes jury)...")
    add_slides_part4(prs)
    print(f"   Slides generees : {len(prs.slides)}")

    output = "Soutenance_PFE_Adrien_TRAVAILLE.pptx"
    prs.save(output)
    print("=" * 70)
    print(f"SUCCES : '{output}' sauvegarde ({len(prs.slides)} diapositives).")
    print("=" * 70)
    return prs


if __name__ == "__main__":
    generate_complete_defense_deck()
