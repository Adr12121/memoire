import sys
from pptx import Presentation

prs = Presentation('Presentation_PFE_180s.pptx')
print(f"Total slides: {len(prs.slides)}")

for i, slide in enumerate(prs.slides):
    print(f"\n--- SLIDE {i+1} ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    # Print without unicode crashes
                    clean_txt = txt.encode('ascii', errors='replace').decode('ascii')
                    print(f"  [Text] {clean_txt[:80]}")
