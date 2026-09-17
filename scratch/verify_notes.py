# -*- coding: utf-8 -*-
from pptx import Presentation

prs = Presentation('Soutenance_PFE_Adrien_TRAVAILLE.pptx')
print(f'Total slides in pptx: {len(prs.slides)}')

for i in [1, 2, 7, 18, 20, 23, 26, 31, 34, 36, 37]:
    slide = prs.slides[i - 1]
    notes = slide.notes_slide.notes_text_frame.text
    print(f'\n--- Slide {i} Notes (length: {len(notes)}) ---')
    for line in notes.splitlines()[:5]:
        print(' ', line)
