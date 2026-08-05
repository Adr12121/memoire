import docx

def update_word_document():
    doc_path = "C:/Users/Topo_4/Documents/AT_PFE/TRAVAILLE_Resume_PFE_JourneesTopo_v2.docx"
    doc = docx.Document(doc_path)
    
    # Textes de remplacement (sans les sauts de ligne cette fois)
    texte_1 = "Le flux de traitement débute par la réception des fichiers bruts (scans PDF ou images) et leur catégorisation automatique. Cette étape est structurante car l'hétérogénéité des archives foncières (plans modernes, livrets historiques de l'ancien cadastre, croquis de bornage) exige des stratégies d'extraction radicalement différentes. Un algorithme de tri (via plan_classifier.py) analyse la géométrie globale du document pour identifier sa typologie. Ce premier routage garantit que chaque archive est orientée vers le sous-système d'analyse visuelle le plus pertinent."
    texte_2 = "Une fois le document classifié, l'extraction de l'information s'opère d'abord par une approche spatiale. Plutôt que de chercher des chaînes de caractères à l'aveugle, le système emploie le réseau de neurones convolutifs YOLOv8 [Jocher et al., 2023] pour localiser visuellement les zones d'intérêt (cartouches, tableaux de parcelles, blocs de signatures). Ce choix technologique se justifie par la robustesse de YOLOv8 face à la dégradation des scans et son excellente rapidité d'inférence. Le recours à des modèles d'analyse documentaire plus complexes, tels que LayoutLM [Bajrami et al., 2023], a été sciemment écarté. Bien que performants, ces modèles exigent des ressources matérielles (calcul GPU intensif) incompatibles avec le parc informatique standard d'un cabinet de géomètre."
    texte_3 = "Après le découpage des zones cibles, le système extrait le contenu textuel (semantic_ocr_engine.py). Les archives mêlant textes dactylographiés et annotations manuscrites, une approche hybride a été implémentée. Pour les textes imprimés et structurés, un moteur de reconnaissance optique de caractères (OCR) déterministe comme Tesseract est privilégié pour sa fiabilité éprouvée [AlKendi et al., 2024]. En revanche, face aux écritures manuscrites complexes, le système bascule dynamiquement sur un Modèle de Vision-Langage (VLM), exécuté localement (Ollama/LLaVA). L'intégration de ces modèles permet d'interpréter le texte tant visuellement que sémantiquement, offrant une transcription robuste là où un OCR classique échoue [Tarride et al., 2024]."
    texte_4 = "L'extraction par intelligence artificielle étant probabiliste, une phase de fiabilisation stricte est requise (coherence_checker.py). Ce module exploite un grand modèle de langage (LLM local) pour vérifier la logique métier des données extraites. Le système croise automatiquement les informations avec des bases de données de référence : validation des communes via le répertoire de l'INSEE et authentification des praticiens via l'annuaire officiel des Géomètres-Experts. À l'issue de cette validation multicritère, un score de confiance quantifié est attribué à l'ensemble du dossier [Sang, 2024]."
    texte_5 = "Conformément aux règles déontologiques de la profession, l'intelligence artificielle n'agit que comme un assistant ; la décision finale appartient à l'expert [Hours, 2023]. Les dossiers pré-structurés sont présentés à l'opérateur via une interface visuelle interactive développée sous Streamlit (app_validation.py). Le géomètre peut y vérifier visuellement la concordance entre le document d'origine et la donnée extraite, puis valider ou corriger les champs présentant un faible score de confiance. Ces corrections humaines sont par la suite utilisées pour réentraîner et affiner les modèles."
    texte_6 = "Une fois les métadonnées consolidées et validées par le professionnel, le système procède à l'association finale des données avec les codes d'opération normés (division, bornage, etc.). Le dossier, parfaitement structuré (geofoncier_api.py), est alors téléversé de manière automatisée sur le portail national via l'interface de programmation (API) de Géofoncier. Cette ultime étape clôture le processus de dématérialisation et d'archivage sécurisé."

    # On va rechercher les paragraphes de l'ancien document
    for i, p in enumerate(doc.paragraphs):
        if "1 Analyse visuelle et localisation" in p.text:
            p.text = ""
            run = p.add_run("1. Classification des documents entrants\n")
            run.bold = True
            p.add_run(texte_1 + "\n\n")
            run2 = p.add_run("2. Détection spatiale et segmentation\n")
            run2.bold = True
            p.add_run(texte_2)
            
        elif "2 Transcription textuelle hybride" in p.text:
            p.text = ""
            run = p.add_run("3. Transcription textuelle par lecture hybride\n")
            run.bold = True
            p.add_run(texte_3)
            
        elif "Fiabilisation, contrôle et intégration" in p.text:
            p.text = ""
            run = p.add_run("4. Contrôle de cohérence sémantique\n")
            run.bold = True
            p.add_run(texte_4 + "\n\n")
            run2 = p.add_run("5. Validation manuelle (Human-in-the-Loop)\n")
            run2.bold = True
            p.add_run(texte_5 + "\n\n")
            run3 = p.add_run("6. Export et interconnexion avec Géofoncier\n")
            run3.bold = True
            p.add_run(texte_6)

    # Sauvegarde dans une V3 pour protéger la V2
    out_path = "C:/Users/Topo_4/Documents/AT_PFE/TRAVAILLE_Resume_PFE_JourneesTopo_v3.docx"
    doc.save(out_path)
    print(f"Fichier principal mis à jour et sauvegardé sous : {out_path}")

if __name__ == "__main__":
    update_word_document()
