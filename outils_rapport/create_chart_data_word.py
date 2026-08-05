from docx import Document

doc = Document()
doc.add_heading('Données pour le diagramme circulaire', 0)

doc.add_paragraph(
    "Voici les données préparées. Pour générer un graphique 100% modifiable directement dans Word :\n"
    "1. Dans votre rapport PFE, allez dans l'onglet Insertion > Graphique > Secteur (ou Anneau).\n"
    "2. Une petite fenêtre Excel va s'ouvrir.\n"
    "3. Copiez/collez simplement le tableau ci-dessous dans cette fenêtre Excel.\n"
    "4. Le graphique se mettra à jour tout seul et vous pourrez modifier les couleurs et les dates à volonté !"
)

table = doc.add_table(rows=6, cols=2)
table.style = 'Table Grid'

# En-têtes
table.rows[0].cells[0].text = 'Prédécesseur'
table.rows[0].cells[1].text = 'Nombre de dossiers (estimatif)'

# Données
data = [
    ('Géomètre A', 6600),
    ('Géomètre B', 6000),
    ('Géomètre C', 4500),
    ('Géomètre D', 3500),
    ('Géomètre E', 3000),
]

for i, (name, val) in enumerate(data):
    table.rows[i+1].cells[0].text = name
    table.rows[i+1].cells[1].text = str(val)

output_path = r'c:\Users\Topo_4\Documents\AT_PFE\Donnees_Graphique_Archives.docx'
doc.save(output_path)
print(f"Document créé : {output_path}")
