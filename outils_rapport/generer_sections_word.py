import docx

doc = docx.Document()

# --- Section 2 ---
doc.add_heading("2. État de l'art et limites des approches classiques : Le choix de l'Intelligence Artificielle", level=1)

doc.add_paragraph("L'extraction classique de texte, basée sur la lecture directe du PDF ou l'utilisation d'expressions régulières sur des gabarits fixes, constitue la méthode traditionnelle d'analyse documentaire. Toutefois, elle est insuffisante pour traiter ce type d'archives. Cette méthode ne fonctionne que sur des documents récents, nativement numériques et sans dégradation. Dès qu'un plan présente une mention manuscrite, une écriture altérée ou une qualité de numérisation moins bonne, les résultats classiques sont inutilisables. Or, ces défauts sont omniprésents dans les documents de l’étude.")

doc.add_paragraph("Face à ces limites, l'application de l'intelligence artificielle pour l'analyse de documents complexes s'impose comme le nouvel état de l'art. Des outils automatisés sont déjà exploités avec succès dans le domaine juridique (Ross, Predictice) et par les notaires (Intellig'IA) [Barthe, 2017]. Dans le domaine de la topographie, l'apprentissage profond (Deep Learning) fait également l'objet de recherches pour l'analyse de nuages de points 3D et la structuration de données spatiales [Poux, 2020]. Ces travaux justifient l'étude de ces algorithmes pour le traitement automatisé des archives foncières.")

doc.add_paragraph("L'évolution des réseaux de neurones convolutifs (CNN) [LeCun et al., 2015] permet en effet un traitement direct des documents en suivant une logique de reconnaissance visuelle (position précise et motif défini) plutôt que textuelle. Dans ce projet, le choix technologique s'est porté sur le modèle de détection d'objets YOLO (You Only Look Once) [Jiang et al., 2022]. Au lieu de rechercher des chaînes de caractères instables, ce modèle identifie les coordonnées spatiales des éléments clés sur la page (cartouches, tableaux de données, signatures), indépendamment de la qualité de la numérisation. La Fig. 1 illustre cette capacité : l'intelligence artificielle parvient à localiser et lire des informations complexes directement sur le plan scanné.")

p_img = doc.add_paragraph()
r_img = p_img.add_run("[ INSÉRER LA FIGURE 1 (Photo avec les étiquettes vertes) ICI ]")
r_img.bold = True

doc.add_paragraph("Enfin, l'analyse numérique de ces archives impose le respect strict des règles de confidentialité. Ces documents contiennent des données à caractère personnel (noms de propriétaires, adresses, historiques de parcelles) sujettes au secret professionnel, tel que défini par le Code des devoirs professionnels de l'Ordre des Géomètres-Experts. Afin de respecter cette norme déontologique et de garantir la conformité au Règlement Général sur la Protection des Données (RGPD), l'outil développé repose sur une architecture d'exécution intégralement locale. L'ensemble des modèles d'intelligence artificielle fonctionne directement sur le matériel informatique du cabinet. Aucun transfert n'est effectué vers des serveurs externes où l'information risquerait d'être rendue publique. Cette protection (Privacy by Design) a dicté l'architecture de la chaîne de traitement, de l'extraction locale jusqu'à la validation par l’opérateur et au versement final sécurisé sur l'API Géofoncier.")


# --- Section 3 ---
doc.add_heading("3. Architecture du pipeline de traitement", level=1)

doc.add_paragraph("Afin de répondre aux contraintes du métier, le système développé repose sur un pipeline automatisé découpé en six étapes successives (voir Fig. 2) :")

p1 = doc.add_paragraph()
p1.add_run("1. Classification des documents.\n").bold = True
p1.add_run("Le système commence par réceptionner les fichiers bruts (scans PDF ou images) et identifie leur typologie (plan récent, livret historique, croquis). Ce premier filtre permet d'orienter immédiatement le document vers les outils d'analyse les plus adaptés à sa structure.")

p2 = doc.add_paragraph()
p2.add_run("2. Détection spatiale.\n").bold = True
p2.add_run("L'algorithme de détection d'objets YOLOv8 [Jiang et al., 2022] analyse ensuite l'image pour localiser avec précision les zones d'intérêt comme les cartouches, les tableaux de données ou les signatures. L'utilisation de modèles d'analyse documentaire plus lourds, comme LayoutLM [Xu et al., 2020], a été exclue ici, car leurs exigences matérielles étaient incompatibles avec les ordinateurs classiques d'un cabinet.")

p3 = doc.add_paragraph()
p3.add_run("3. Lecture hybride.\n").bold = True
p3.add_run("Une fois les zones découpées, le programme extrait le texte. Pour garantir des résultats exploitables sur des documents très hétérogènes, le système choisit automatiquement son outil de lecture. Un moteur de reconnaissance optique classique (Tesseract) est utilisé pour les textes imprimés standards. En revanche, pour déchiffrer les écritures manuscrites complexes, le système bascule sur un Modèle de Vision-Langage exécuté localement (Ollama/LLaVA) [Tarride et al., 2024].")

p4 = doc.add_paragraph()
p4.add_run("4. Contrôle de cohérence.\n").bold = True
p4.add_run("Les informations brutes obtenues sont confrontées à un algorithme de vérification. Ce module utilise un grand modèle de langage (Ollama LLM) pour vérifier la logique des données, croiser les noms avec la base de l'Ordre des géomètres, et valider la géographie avec les bases de l'INSEE. Un score de confiance global est alors calculé et attribué au dossier [Sang, 2024].")

p5 = doc.add_paragraph()
p5.add_run("5. Validation manuelle.\n").bold = True
p5.add_run("Avant l'envoi définitif, le système impose une étape de validation humaine (Human-in-the-loop). Une interface visuelle de contrôle (développée avec Streamlit) permet à l'opérateur de revoir les dossiers, de corriger les informations incertaines et d'entraîner le modèle sur ses erreurs via l'apprentissage de ces corrections.")

p6 = doc.add_paragraph()
p6.add_run("6. Export Géofoncier.\n").bold = True
p6.add_run("Une fois le dossier validé et parfaitement structuré, le système associe les codes d'opération définitifs. Les données sont alors soumises et téléversées automatiquement sur le portail officiel via l'interface de programmation (API) Géofoncier, clôturant le processus d'archivage.")

doc.save("C:/Users/Topo_4/Documents/AT_PFE/Sections_Reecrites.docx")
print("Fichier Sections_Reecrites.docx généré avec succès.")
