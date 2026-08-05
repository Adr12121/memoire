import openpyxl
from openpyxl.chart import DoughnutChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, Alignment, PatternFill

# Création du classeur
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Données"

# Données Anonymisées
data = [
    ['Cabinet', 'Dossiers'],
    ['Cabinet A', 6600],
    ['Cabinet B', 6000],
    ['Cabinet C', 4500],
    ['Cabinet D', 3500],
    ['Cabinet E', 3000],
]

# Ajout des données dans la feuille
for row in data:
    ws.append(row)

# Mise en forme du tableau de données
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

# Ajustement de la largeur des colonnes
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 12

# --- Création du Graphique en Anneau (Doughnut Chart) ---
chart = DoughnutChart()
chart.title = "RÉPARTITION DU PASSIF DOCUMENTAIRE"
chart.style = 2  # Style moderne par défaut
# chart.holeSize = 50 # openpyxl DoughnutChart a holeSize

# Sélection des données
labels = Reference(ws, min_col=1, min_row=2, max_row=6)
values = Reference(ws, min_col=2, min_row=1, max_row=6)

chart.add_data(values, titles_from_data=True)
chart.set_categories(labels)

# Ajout d'étiquettes de données (pourcentages)
chart.dataLabels = DataLabelList()
chart.dataLabels.showPercent = True
chart.dataLabels.showVal = False
chart.dataLabels.showCatName = False
chart.dataLabels.showLegendKey = False

# Positionnement du graphique
ws.add_chart(chart, "D2")

# Sauvegarde
output_path = r"c:\Users\Topo_4\Documents\AT_PFE\Repartition_Geometres_Modifiable.xlsx"
wb.save(output_path)
print(f"FICHIER EXCEL MODIFIABLE GENERE : {output_path}")
